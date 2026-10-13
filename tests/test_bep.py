"""Smoke tests for the BEP Builder CLI, checkers, and app parity.

Run:  python tests/test_bep.py      (or: pytest tests/)
Needs openpyxl for the workbook round trip; Node.js (optional) for app/CLI parity.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CLI = ROOT / "skills/bep-builder-cli/scripts"
NCS = ROOT / "skills/national-cad-standard/scripts"
NUM = ROOT / "skills/ios-numbers-spreadsheets/scripts"
sys.path[:0] = [str(CLI), str(NCS)]

import iso13567  # noqa: E402
import check_names  # noqa: E402


def bep(*args, cwd=None):
    r = subprocess.run([sys.executable, str(CLI / "bep.py"), *map(str, args)], capture_output=True, text=True, cwd=cwd)
    return r.returncode, r.stdout + r.stderr


def test_ncs_names():
    assert check_names.check_layer("A-WALL-FULL") == []
    assert check_names.check_layer("A-WALL-D") == []
    assert check_names.check_layer("a-wall")
    assert check_names.check_layer("A_WALL")
    assert check_names.check_layer("A-WALL-FULL-Q")
    assert check_names.check_layer("A-WALL-2", "status")
    assert check_names.check_sheet("A-101") == [] and check_names.check_sheet("AD101") == []
    assert check_names.check_sheet("S501") and check_names.check_sheet("A-100") and check_names.check_sheet("AD-101")


def test_iso13567_profile():
    p = iso13567.Profile(1, "element", "status")
    assert iso13567.check("A-WALL-FULL-TEXT-N", p) == []
    assert iso13567.check("A-WALL-____-____-N", p) == []
    assert iso13567.check("G-ANNO-TEXT-N", p)            # ANNO banned
    assert iso13567.check("A-TEXT-____-____-N", p)       # annotation code as element
    assert iso13567.check("A-WALL-FULL-TEXT-2", p)       # phase in a status project
    assert iso13567.convert("G-ANNO-TTLB-N", p, None, "SHET")[0] == "G-SHET-____-TTLB-N"
    assert iso13567.convert("A-ANNO-DIMS", p, None, "SHET")[0] is None   # element annotation needs a person
    assert iso13567.convert("A-WALL-FULL-D", p)[0] == "A-WALL-FULL-____-D"
    v = iso13567.Profile(1, "view", "phase")
    assert iso13567.check("A-WALL-ELEV-TEXT-2", v) == []
    assert iso13567.check("A-WALL-FULL-TEXT-2", v)


def test_cli_plan_lifecycle():
    d = Path(tempfile.mkdtemp())
    try:
        assert bep("init", d / "p.json", "--agency", "usace")[0] == 0
        assert bep("set", d / "p.json", "project.cui", "maybe")[0] != 0       # validated
        assert bep("set", d / "p.json", "project.nmae", "X")[0] != 0          # unknown field
        assert bep("set", d / "p.json", "standards.iso13567", "maybe")[0] != 0  # not yes/no
        assert bep("show", d / "p.json", "project.zzz")[0] != 0
        assert bep("remove", d / "p.json", "team", "-1")[0] != 0
        assert bep("qc", d / "p.json", "visual", "--off")[0] != 0
        assert bep("qc", d / "p.json", "a0", "--tool", "Revit")[0] != 0
        assert bep("use", d / "p.json", "coord", "--party", "BIM manager")[0] == 0
        code, out = bep("validate", d / "p.json", "--strict")
        assert code == 1 and "TODO items" in out
        code, out = bep("export", d / "p.json")
        assert code == 0 and "USACE / DoD overrides" in out and "```mermaid" in out
        plan = json.loads((d / "p.json").read_text())
        assert plan["standards"]["sheetCode"] == "SHET" and plan["standards"]["iso13567"] is False
    finally:
        shutil.rmtree(d)


def test_process_map_example():
    code, out = bep("process-map", ROOT / "examples/example-plan.json")
    assert code == 0
    assert 'u_authoring -->|"Coordination model · LOD 300 · Weekly"| u_coord' in out
    assert "stub_construction" in out                                  # skipped phase gets a placeholder
    assert 'stub_construction -->|"COBie handover · LOD 500 data · Closeout"| u_cobie' in out


def test_examples_in_sync():
    """examples/example-plan.json, 'bep init --example' and the app's example are the same plan."""
    sys.path.insert(0, str(CLI))
    import bep as cli
    ref = cli.normalize(json.loads((ROOT / "examples/example-plan.json").read_text()))
    assert cli.normalize(cli.example_plan()) == ref
    assert not cli.collect_todos(cli.build_doc(ref))
    node = shutil.which("node")
    if not node:
        print("  (skipped app check: Node.js not installed)")
        return
    html = (ROOT / "app/index.html").read_text()
    js = html.split("<script>")[1].split("</script>")[0].split("/* ---------- Events ---------- */")[0]
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False) as fh:
        fh.write(js + "\nconsole.log(JSON.stringify(normalize(examplePlan())));\n")
    try:
        r = subprocess.run([node, fh.name], capture_output=True, text=True)
        assert r.returncode == 0, r.stderr
    finally:
        os.unlink(fh.name)
    assert cli.normalize(json.loads(r.stdout)) == ref


def test_maps_workbook_round_trip():
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  (skipped: openpyxl not installed)")
        return
    d = Path(tempfile.mkdtemp())
    try:
        shutil.copy(ROOT / "examples/example-plan.json", d / "p.json")
        assert bep("maps-workbook", d / "p.json", "-o", d / "maps.xlsx")[0] == 0
        code, out = bep("apply-maps", d / "p.json", d / "maps.xlsx")
        assert code == 0 and "No changes" in out                       # untouched workbook = no-op
        wb = load_workbook(d / "maps.xlsx")
        ws = wb["Uses"]
        for row in ws.iter_rows(min_row=2):
            if row[2].value == "Phase planning (4D)":
                row[4].value, row[5].value = "Y", "Contractor VDC lead"
        wb.save(d / "edited.xlsx")
        code, out = bep("apply-maps", d / "p.json", d / "edited.xlsx")
        assert code == 0 and "Selected Phase planning (4D)" in out
        plan = json.loads((d / "p.json").read_text())
        assert plan["uses"]["4d"]["on"] and plan["uses"]["4d"]["party"] == "Contractor VDC lead"
        assert "No changes" in bep("apply-maps", d / "p.json", d / "edited.xlsx")[1]   # idempotent
        diff = subprocess.run([sys.executable, str(NUM / "diff_tables.py"), d / "maps.xlsx", d / "edited.xlsx",
                               "--key", "Uses=Project,BIM use"], capture_output=True, text=True).stdout
        assert "Phase planning (4D) | Selected (Y/N) | N | Y" in diff
    finally:
        shutil.rmtree(d)


def test_app_matches_cli():
    node = shutil.which("node")
    if not node:
        print("  (skipped: Node.js not installed)")
        return
    html = (ROOT / "app/index.html").read_text()
    js = html.split("<script>")[1].split("</script>")[0].split("/* ---------- Events ---------- */")[0]
    plan = json.loads((ROOT / "examples/example-plan.json").read_text())
    names = ["A-WALL-FULL-TEXT-N", "G-ANNO-TEXT-N", "A-TEXT-____-____-N", "G-ANNO-TTLB-N", "A-ANNO-DIMS", "A-WALL-FULL-D"]
    harness = js + f"""
plan = normalize({json.dumps(plan)});
const p = profileFor();
const out = {{ map: mapToMermaid(buildMap()), iso: {json.dumps(names)}.map(n => [isoCheck(n, p).length === 0, isoConvert(n, p, null, "SHET")[0]]) }};
console.log(JSON.stringify(out));
"""
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False) as fh:
        fh.write(harness)
    try:
        r = subprocess.run([node, fh.name], capture_output=True, text=True)
        assert r.returncode == 0, r.stderr
        app = json.loads(r.stdout)
    finally:
        os.unlink(fh.name)
    assert app["map"] == bep("process-map", ROOT / "examples/example-plan.json")[1]
    prof = iso13567.Profile(1, "element", "status")
    assert app["iso"] == [[iso13567.check(n, prof) == [], iso13567.convert(n, prof, None, "SHET")[0]] for n in names]


# ---------------------------------------------------------------- regression tests for the
# high-priority review items (CLI data safety, workbook injection, app saving and downloads)
def _app_js():
    html = (ROOT / "app/index.html").read_text()
    return html.split("<script>")[1].split("</script>")[0].split("/* ---------- Events ---------- */")[0]


# Just enough DOM for the app's helpers to run under Node.
DOM_STUB = """
globalThis.__clicked = [];
globalThis.__stored = {};
Object.defineProperty(globalThis, "localStorage", { configurable: true,
  value: { setItem: (k, v) => { __stored[k] = v; }, getItem: k => __stored[k] ?? null } });
globalThis.document = {
  querySelector: () => ({ style: {} }),
  body: { appendChild() {} },
  createElement: () => ({ style: {}, click() { __clicked.push({ href: this.href, download: this.download }); }, remove() {} })
};
"""


def _node(code):
    node = shutil.which("node")
    if not node:
        return None
    with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False) as fh:
        fh.write(DOM_STUB + _app_js() + "\n" + code)
    try:
        r = subprocess.run([node, fh.name], capture_output=True, text=True)
        assert r.returncode == 0, r.stderr
        return json.loads(r.stdout)
    finally:
        os.unlink(fh.name)


def test_set_refuses_sections_and_keeps_file_valid():
    d = Path(tempfile.mkdtemp())
    try:
        assert bep("init", d / "p.json")[0] == 0
        before = (d / "p.json").read_text()
        for path, value in (("project", "foo"), ("team", "foo"), ("uses.coord", "x"), ("team.-1.name", "Bob")):
            code, out = bep("set", d / "p.json", path, value)
            assert code != 0 and "Couldn't set" in out, out
        assert (d / "p.json").read_text() == before                    # nothing was written
        assert bep("status", d / "p.json")[0] == 0                      # file still loads
        assert bep("set", d / "p.json", "team.0.name", "Ana")[0] == 0
        assert json.loads((d / "p.json").read_text())["team"][0]["name"] == "Ana"
        assert [f.name for f in d.iterdir()] == ["p.json"]              # atomic save leaves no temp file
        if os.name == "posix":
            os.chmod(d / "p.json", 0o700)   # owner-only, but not mkstemp's 0o600
            assert bep("set", d / "p.json", "project.name", "X")[0] == 0
            assert (d / "p.json").stat().st_mode & 0o777 == 0o700       # permissions kept
            (d / "link.json").symlink_to("p.json")
            assert bep("set", d / "link.json", "project.name", "Via link")[0] == 0
            assert (d / "link.json").is_symlink()                       # link kept
            assert json.loads((d / "p.json").read_text())["project"]["name"] == "Via link"
            if os.geteuid() != 0:                                       # root ignores file modes
                os.chmod(d / "p.json", 0o400)
                code, out = bep("set", d / "p.json", "project.name", "Y")
                assert code != 0 and "read-only" in out, out
                assert json.loads((d / "p.json").read_text())["project"]["name"] == "Via link"
                os.chmod(d / "p.json", 0o600)
    finally:
        shutil.rmtree(d)


def test_partial_rows_are_filled_from_templates():
    sys.path.insert(0, str(CLI))
    import bep as cli
    plan = json.loads((ROOT / "examples/example-plan.json").read_text())
    del plan["team"][0]["email"]
    del plan["exchanges"][0]["lod"]
    plan["standards"]["codes"] = [{"code": "ACOU"}]
    plan["collab"] = "damaged"                                         # a section saved as a scalar
    plan["goals"].append("damaged")                                    # a row saved as a scalar
    d = Path(tempfile.mkdtemp())
    try:
        (d / "p.json").write_text(json.dumps(plan))
        code, out = bep("export", d / "p.json", "-o", d / "p.md")
        assert code == 0, out
        n = cli.normalize(plan)
        assert n["team"][0]["email"] == "" and n["exchanges"][0]["lod"] == ""
        assert n["standards"]["codes"] == [{"code": "ACOU", "meaning": ""}]
        assert n["collab"] == cli.blank_plan()["collab"]
        assert n["goals"][-1] == {"goal": "", "uses": ""}
        (d / "s.json").write_text(json.dumps({"standards": "damaged", "team": ["foo"],
                                              "exchanges": [{"what": None, "lod": None}]}))
        for cmd in ("status", "export", "validate", "process-map"):
            code, out = bep(cmd, d / "s.json")
            assert "Traceback" not in out, out
        assert "TODO: Coordination model: LOD" in (d / "p.md").read_text()
    finally:
        shutil.rmtree(d)


def test_maps_workbook_writes_plan_text_as_text():
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  (skipped: openpyxl not installed)")
        return
    evil = '=HYPERLINK("http://example.invalid","Click")'
    d = Path(tempfile.mkdtemp())
    try:
        shutil.copy(ROOT / "examples/example-plan.json", d / "p.json")
        assert bep("use", d / "p.json", "coord", "--party", evil)[0] == 0
        assert bep("add", d / "p.json", "exchanges", f"what={evil}", "lod=LOD 300")[0] == 0
        assert bep("maps-workbook", d / "p.json", "-o", d / "maps.xlsx")[0] == 0
        wb = load_workbook(d / "maps.xlsx")
        cells = [c for ws in (wb["Uses"], wb["Exchanges"]) for row in ws.iter_rows(min_row=2) for c in row
                 if c.value == evil]
        assert len(cells) == 2 and all(c.data_type == "s" for c in cells)   # text, not formulas
        code, out = bep("apply-maps", d / "p.json", d / "maps.xlsx")
        assert code == 0 and "No changes" in out, out                       # round trip unchanged
    finally:
        shutil.rmtree(d)


def test_app_downloads_without_the_claude_capability():
    out = _node("""
(async () => {
  const ok = await saveFile("plan-bep.md", "# Plan", "text/markdown");
  console.log(JSON.stringify({ ok, clicked: __clicked }));
})();
""")
    if out is None:
        print("  (skipped: Node.js not installed)")
        return
    assert out["ok"] is True
    assert len(out["clicked"]) == 1 and out["clicked"][0]["download"] == "plan-bep.md"
    assert out["clicked"][0]["href"].startswith("blob:")


def test_app_shared_save_keeps_teammates_changes():
    out = _node("""
const merge = (a, b) => { for (const k in b) { const x = a[k], y = b[k];
  if (y && typeof y === "object" && !Array.isArray(y) && x && typeof x === "object" && !Array.isArray(x)) merge(x, y); else a[k] = y; } return a; };
(async () => {
  plan = normalize(examplePlan());
  const server = { doc: { plan: clone(plan) }, sets: 0, updates: 0 };
  planRef = {   // db document: set replaces; update merges objects, replaces arrays, needs the doc
    set: async d => { server.sets++; server.doc = clone(d); },
    update: async d => { server.updates++; if (!server.doc) throw { code: "invalid_argument" }; merge(server.doc, clone(d)); }
  };
  mode = "shared"; docExists = true;
  server.doc.plan.collab.cde = "Teammate CDE";          // a teammate saved first
  plan.project.name = "Mine"; scheduleSave("project.name");
  plan.team[1].name = "New name"; scheduleSave("team.1.name");
  clearTimeout(saveTimer); await flush();
  const first = { plan: clone(server.doc.plan), sets: server.sets, updates: server.updates };

  server.doc = null;                                     // the document was deleted
  plan.project.owner = "Owner X"; scheduleSave("project.owner");
  clearTimeout(saveTimer); await flush();
  const second = { owner: server.doc.plan.project.owner, name: server.doc.plan.project.name, sets: server.sets };

  const remote = normalize(examplePlan()); remote.goals[0].goal = "Teammate goal";
  const local = normalize(examplePlan()); local.project.location = "My unsaved edit";
  const merged = overlayEdits(remote, local, ["project.location"]);

  const patches = {
    name: buildPatch(normalize(blankPlan()), ["project.name"]),
    statusField: buildPatch(normalize(blankPlan()), ["standards.statusField"])
  };

  // An edit whose save is still in flight survives "Load their changes".
  plan = normalize(examplePlan()); fullSave = false; dirty.clear(); docExists = true;
  let release; const gate = new Promise(r => release = r);
  planRef = { set: async () => {}, update: async () => { await gate; } };
  plan.project.location = "Saving now"; scheduleSave("project.location");
  clearTimeout(saveTimer); const pending = flush();
  applyRemote({ plan: remote, updatedBy: "teammate", updatedAt: new Date().toISOString() });
  const inFlightKept = plan.project.location === "Saving now" && plan.goals[0].goal === "Teammate goal";
  release(); await pending;

  // A saved edit whose echo hasn't arrived yet survives "Load their changes".
  planRef = { set: async () => {}, update: async () => {} };
  plan = normalize(examplePlan()); dirty.clear(); fullSave = false; docExists = true; committed = [];
  plan.project.owner = "Just saved"; scheduleSave("project.owner");
  clearTimeout(saveTimer); await flush();
  applyRemote({ plan: remote, updatedBy: "teammate", updatedAt: new Date().toISOString() });
  const committedKept = plan.project.owner === "Just saved" && plan.goals[0].goal === "Teammate goal";
  const echoPlan = clone(plan);
  onRemote({ exists: true, data: () => ({ plan: echoPlan, clientId, updatedAt: committedAt }) });
  const committedCleared = committed.length === 0;

  // A whole-plan save clears fields kept from an earlier partial save.
  plan.project.owner = "Partial"; scheduleSave("project.owner");
  clearTimeout(saveTimer); await flush();
  const hadCommitted = committed.length > 0;
  scheduleSave(); clearTimeout(saveTimer); await flush();       // e.g. Load example / Clear plan
  const fullClears = hadCommitted && committed.length === 0 && committedAt === null;

  // A pending whole-plan change (Clear plan) survives "Load their changes".
  plan = blankPlan(); fullSave = true; dirty.clear();
  applyRemote({ plan: remote, updatedBy: "teammate", updatedAt: new Date().toISOString() });
  const keptCleared = plan.project.name === "" && plan.goals[0].goal === "";

  // Saving before the first snapshot: check whether the shared plan exists, then merge.
  const slow = { doc: { plan: clone(remote) }, sets: 0, updates: 0, gets: 0 };
  planRef = {
    get: async () => { slow.gets++; return { exists: !!slow.doc }; },
    set: async d => { slow.sets++; slow.doc = clone(d); },
    update: async d => { slow.updates++; merge(slow.doc, clone(d)); }
  };
  plan = normalize(examplePlan()); fullSave = false; dirty.clear(); docExists = null;
  plan.project.phase = "Construction"; scheduleSave("project.phase");
  clearTimeout(saveTimer); await flush();
  const unknownFirst = { gets: slow.gets, sets: slow.sets, updates: slow.updates,
    keptGoal: slow.doc.plan.goals[0].goal, phase: slow.doc.plan.project.phase };

  // An edit made while connecting is saved once the first snapshot arrives.
  const conn = { doc: { plan: clone(remote) }, updates: 0 };
  planRef = { get: async () => ({ exists: true }), set: async () => {},
              update: async d => { conn.updates++; merge(conn.doc, clone(d)); } };
  plan = normalize(blankPlan()); dirty.clear(); fullSave = false; docExists = null; mode = "connecting";
  plan.project.owner = "Typed while connecting"; scheduleSave("project.owner");
  const queuedWhileConnecting = !!saveTimer;
  onRemote({ exists: true, data: () => ({ plan: clone(remote), clientId: "teammate" }) });
  const scheduledOnConnect = !!saveTimer;
  clearTimeout(saveTimer); await flush();
  const connecting = { queuedWhileConnecting, scheduledOnConnect, updates: conn.updates,
    owner: conn.doc.plan.project.owner, keptGoal: conn.doc.plan.goals[0].goal };

  // A retry that fires after falling back to browser-only saving saves locally, not remotely.
  let remoteCalls = 0;
  planRef = { get: async () => { remoteCalls++; return { exists: true }; }, set: async () => { remoteCalls++; }, update: async () => { remoteCalls++; } };
  mode = "shared"; docExists = true; dirty.clear(); fullSave = false;
  plan = normalize(examplePlan()); plan.project.location = "Edited during outage"; dirty.add("project.location");
  mode = "local";                                      // the snapshot listener reported the store gone
  await flush();
  const localFallback = { remoteCalls, dirty: dirty.size,
    stored: JSON.parse(__stored[LS_KEY] || "{}").project?.location || null };
  mode = "shared";

  // Falling back to browser-only saving drops a teammate version waiting behind the banner.
  mode = "shared"; latestRemote = { plan: remote }; pendingRemote = true;
  goLocal();
  const localDropsStale = latestRemote === null && pendingRemote === false && mode === "local";
  mode = "shared";

  // A deleted shared plan clears a teammate version still waiting behind the banner.
  latestRemote = { plan: remote }; pendingRemote = true;
  onRemote({ exists: false });
  const staleCleared = latestRemote === null && pendingRemote === false;
  console.log(JSON.stringify({ first, second,
    roots: ["team.1.name", "standards.codes.0.code", "uses.coord.on", "project.name"].map(patchRoot),
    merged: { goal: merged.goals[0].goal, location: merged.project.location }, keptCleared, inFlightKept, staleCleared, unknownFirst, connecting, patches, localFallback, committedKept, committedCleared, localDropsStale, fullClears }));
})();
""")
    if out is None:
        print("  (skipped: Node.js not installed)")
        return
    f = out["first"]
    assert f["sets"] == 0 and f["updates"] == 1                     # one merge-write, no whole-plan overwrite
    assert f["plan"]["collab"]["cde"] == "Teammate CDE"              # teammate's change survives
    assert f["plan"]["project"]["name"] == "Mine" and f["plan"]["team"][1]["name"] == "New name"
    s = out["second"]
    assert s["sets"] == 1 and s["owner"] == "Owner X" and s["name"] == "Mine"   # recreated whole
    assert out["roots"] == ["team", "standards.codes", "uses.coord.on", "project.name"]
    assert out["merged"] == {"goal": "Teammate goal", "location": "My unsaved edit"}
    assert out["keptCleared"] is True and out["inFlightKept"] is True and out["staleCleared"] is True
    assert out["unknownFirst"] == {"gets": 1, "sets": 0, "updates": 1,                # merged, not overwritten
                                   "keptGoal": "Teammate goal", "phase": "Construction"}
    assert out["committedKept"] is True and out["committedCleared"] is True and out["localDropsStale"] is True and out["fullClears"] is True
    assert out["localFallback"] == {"remoteCalls": 0, "dirty": 0, "stored": "Edited during outage"}
    assert out["patches"]["name"] == {"project": {"name": ""}}                # no derived status riding along
    assert out["patches"]["statusField"] == {"standards": {"statusField": "status", "status": True}}
    assert out["connecting"] == {"queuedWhileConnecting": False, "scheduledOnConnect": True, "updates": 1,
                                 "owner": "Typed while connecting", "keptGoal": "Teammate goal"}


def test_app_normalize_matches_cli_on_damaged_plans():
    sys.path.insert(0, str(CLI))
    import bep as cli
    damaged = {"standards": "damaged", "collab": "damaged",
               "team": ["foo", {"role": "BIM manager", "email": None}],
               "exchanges": [{"what": None, "lod": "LOD 300"}], "goals": "damaged"}
    out = _node(f"console.log(JSON.stringify(normalize({json.dumps(damaged)})));")
    if out is None:
        print("  (skipped: Node.js not installed)")
        return
    assert out == cli.normalize(json.loads(json.dumps(damaged)))


def test_pages_build_is_a_full_document():
    sys.path.insert(0, str(ROOT / "tools"))
    import build_pages
    out = build_pages.build((ROOT / "app/index.html").read_text())
    head, body = out.split("</head>", 1)
    assert out.startswith("<!doctype html>\n<html lang=\"en\">")
    assert '<meta charset="utf-8">' in head and 'name="viewport"' in head
    assert "<title>BIM Execution Plan Builder</title>" in head and "<style>" in head
    assert "--paper:" in head                                         # the app's stylesheet moved into <head>
    assert body.lstrip().startswith('<body>\n<div class="wrap">') and body.rstrip().endswith("</body>\n</html>")
    assert build_pages.build(out) == out                              # idempotent

if __name__ == "__main__":
    failed = 0
    for name, fn in [(k, v) for k, v in dict(globals()).items() if k.startswith("test_") and callable(v)]:
        try:
            fn()
            print(f"PASS {name}")
        except Exception as e:
            failed += 1
            print(f"FAIL {name}: {type(e).__name__}: {e}")
    sys.exit(1 if failed else 0)
